"""Tests for the M&V report generation package (reporting/).

The authoritative numeric oracle is the XLSX round-trip (class B): the file is
parsed back and its cells are asserted equal to the source dataclass fields. PDF
assertions (class C) are deliberately weaker presence/smoke checks, because PDF
text extraction reflows and splits numbers — do not tighten them into exact
numeric checks.
"""
from __future__ import annotations

import math
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
import pypdf
import pytest
from openpyxl import load_workbook

from engine.baseline import BaselineModel
from engine.mv import SavingsSummary
from reporting import build_mv_report_pdf, build_mv_report_xlsx, format as fmt, verdicts
from reporting.meta import ReportMeta

ROOT = Path(__file__).resolve().parents[1]


# --- Fixtures (build dataclasses directly; no model fitting) -----------------
def make_summary(**kw) -> SavingsSummary:
    d = dict(
        avoided_kwh=164386.0, pct_saving=0.06, inr_saved=1315088.0, tonnes_co2_avoided=116.7,
        days=270, annual_kwh=222226.0, annual_inr=1777808.0, annual_tco2=157.8,
        lag1_autocorr=0.02, n_effective=260.0, t_stat=70.0, p_value=1e-50, is_significant=True,
        avoided_kwh_ci_low=158806.0, avoided_kwh_ci_high=169967.0, ci_confidence=0.90,
    )
    d.update(kw)
    return SavingsSummary(**d)


def make_model(**kw) -> BaselineModel:
    d = dict(
        drivers=["production_tonnes", "ambient_temp_c", "operating_hours"],
        coefficients={"const": 181.5, "production_tonnes": 66.7, "ambient_temp_c": 50.5,
                      "operating_hours": 95.6},
        r_squared=0.922, cv_rmse=0.026, cv_rmse_oos=0.042, n_obs=266, dropped=[],
        baseline_period=("2024-12-05", "2025-08-31"), durbin_watson=1.95, max_vif=1.0,
        bp_pvalue=0.36, model=object(),
    )
    d.update(kw)
    return BaselineModel(**d)


def make_avoided(n: int = 60) -> pd.DataFrame:
    dates = pd.date_range("2025-09-01", periods=n, freq="D")
    avoided = np.full(n, 600.0)
    return pd.DataFrame({
        "date": dates, "actual_kwh": 9300.0, "expected_kwh": 9900.0,
        "avoided_kwh": avoided, "cumulative_avoided_kwh": np.cumsum(avoided),
    })


def make_meta(**kw) -> ReportMeta:
    d = dict(
        facility_name="Test Plant", report_id="WW-1", prepared_by="QA",
        baseline_start="2024-12-05", baseline_end="2025-08-31",
        reporting_start="2025-09-01", reporting_end="2025-10-30",
        tariff_inr_per_kwh=8.0, emission_factor=0.71, generated_at="2026-06-06T00:00:00",
    )
    d.update(kw)
    return ReportMeta(**d)


def _summary_cells(xlsx_bytes: bytes, sheet: str) -> dict:
    wb = load_workbook(BytesIO(xlsx_bytes))
    return {r[0].value: r[1].value for r in wb[sheet].iter_rows() if r[0].value}


# --- A. File validity --------------------------------------------------------
def test_pdf_returns_bytes_with_magic() -> None:
    out = build_mv_report_pdf(make_summary(), make_model(), make_avoided(), make_meta())
    assert isinstance(out, bytes) and len(out) > 1000
    assert out[:5] == b"%PDF-"


def test_xlsx_returns_bytes_with_magic() -> None:
    out = build_mv_report_xlsx(make_summary(), make_model(), make_avoided(), make_meta())
    assert isinstance(out, bytes) and len(out) > 1000
    assert out[:4] == b"PK\x03\x04"


def test_xlsx_has_expected_sheets() -> None:
    out = build_mv_report_xlsx(make_summary(), make_model(), make_avoided(), make_meta())
    wb = load_workbook(BytesIO(out))
    assert set(wb.sheetnames) == {"Summary", "Baseline", "Avoided", "Methodology"}


# --- B. XLSX round-trip (authoritative numeric oracle) -----------------------
def test_xlsx_summary_numbers_match() -> None:
    summary = make_summary()
    cells = _summary_cells(build_mv_report_xlsx(summary, make_model(), make_avoided(), make_meta()), "Summary")
    assert cells["Avoided energy (kWh)"] == pytest.approx(summary.avoided_kwh)
    assert cells["Cost saved (INR)"] == pytest.approx(summary.inr_saved)
    assert cells["CO2 avoided (t)"] == pytest.approx(summary.tonnes_co2_avoided)
    assert cells["Avoided kWh CI low"] == pytest.approx(summary.avoided_kwh_ci_low)
    assert cells["Avoided kWh CI high"] == pytest.approx(summary.avoided_kwh_ci_high)
    assert cells["Significant"] is True


def test_xlsx_baseline_numbers_match() -> None:
    model = make_model()
    cells = _summary_cells(build_mv_report_xlsx(make_summary(), model, make_avoided(), make_meta()), "Baseline")
    assert cells["R-squared"] == pytest.approx(model.r_squared)
    assert cells["CV(RMSE) in-sample"] == pytest.approx(model.cv_rmse)
    assert cells["CV(RMSE) out-of-sample"] == pytest.approx(model.cv_rmse_oos)
    assert cells["Durbin-Watson"] == pytest.approx(model.durbin_watson)
    assert cells["Observations"] == model.n_obs


def test_xlsx_avoided_rows_roundtrip() -> None:
    avoided = make_avoided(60)
    out = build_mv_report_xlsx(make_summary(), make_model(), avoided, make_meta())
    wb = load_workbook(BytesIO(out))
    wsa = wb["Avoided"]
    rows = list(wsa.iter_rows(values_only=True))
    header = rows[0]
    data = rows[1:]
    assert len(data) == len(avoided)
    col = header.index("avoided_kwh")
    assert sum(r[col] for r in data) == pytest.approx(avoided["avoided_kwh"].sum())


def test_xlsx_coefficients_complete() -> None:
    model = make_model()
    wb = load_workbook(BytesIO(build_mv_report_xlsx(make_summary(), model, make_avoided(), make_meta())))
    terms = {r[0].value for r in wb["Baseline"].iter_rows() if r[0].value}
    for k in model.coefficients:
        assert k in terms


def test_xlsx_oos_none_renders_na() -> None:
    cells = _summary_cells(
        build_mv_report_xlsx(make_summary(), make_model(cv_rmse_oos=None), make_avoided(), make_meta()),
        "Baseline",
    )
    assert cells["CV(RMSE) out-of-sample"] == "n/a"


def test_xlsx_max_vif_nan_renders_na() -> None:
    cells = _summary_cells(
        build_mv_report_xlsx(make_summary(), make_model(max_vif=float("nan")), make_avoided(), make_meta()),
        "Baseline",
    )
    assert cells["Max VIF"] == "n/a"


def test_xlsx_methodology_records_assumptions() -> None:
    cells = _summary_cells(
        build_mv_report_xlsx(make_summary(), make_model(), make_avoided(), make_meta(tariff_inr_per_kwh=9.5)),
        "Methodology",
    )
    assert cells["Tariff (INR/kWh)"] == pytest.approx(9.5)
    assert cells["Emission factor (kgCO2/kWh)"] == pytest.approx(0.71)


# --- C. PDF content (presence/smoke only) ------------------------------------
def _pdf_text(pdf_bytes: bytes) -> str:
    reader = pypdf.PdfReader(BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def test_pdf_contains_facility_and_ipmvp() -> None:
    text = _pdf_text(build_mv_report_pdf(make_summary(), make_model(), make_avoided(), make_meta()))
    assert "Test Plant" in text
    assert "IPMVP" in text and "Option C" in text


def test_pdf_contains_key_sections() -> None:
    text = _pdf_text(build_mv_report_pdf(make_summary(), make_model(), make_avoided(), make_meta()))
    for token in ("Baseline", "savings", "CUSUM", "Methodology", "Limitations"):
        assert token in text


def test_pdf_embeds_an_image() -> None:
    reader = pypdf.PdfReader(BytesIO(build_mv_report_pdf(make_summary(), make_model(), make_avoided(), make_meta())))
    has_image = any(
        (page.get("/Resources", {}).get("/XObject")) for page in reader.pages
    )
    assert has_image


# --- D. Formatting (pure) ----------------------------------------------------
def test_format_inr_lakhs() -> None:
    assert fmt.inr_lakhs(1_500_000) == "₹15.00 L"
    assert fmt.inr_lakhs(1_500_000, ascii_only=True) == "INR 15.00 L"


def test_format_mwh_tonnes_pct() -> None:
    assert fmt.mwh(164386) == "164.4 MWh"
    assert fmt.tonnes(1234.4) == "1,234 t"   # avoid .5 (Python rounds half-to-even)
    assert fmt.tonnes(116.7) == "117 t"
    assert fmt.pct(0.083) == "8.3%"


def test_format_coefficient_equation_const_first() -> None:
    eq = fmt.coefficient_equation({"const": 181.5, "production_tonnes": 66.7, "ambient_temp_c": -5.0})
    assert eq.startswith("energy_kwh = 181.5")
    assert "+ 66.7·production_tonnes" in eq
    assert "- 5.0·ambient_temp_c" in eq


def test_format_na_or() -> None:
    assert fmt.na_or(None) == "n/a"
    assert fmt.na_or(float("nan")) == "n/a"
    assert fmt.na_or(1.0, "{:.1f}") == "1.0"


# --- E. ASHRAE verdict logic -------------------------------------------------
def test_ashrae_prefers_oos() -> None:
    val, which = verdicts.ashrae_cvrmse_used(make_model(cv_rmse=0.10, cv_rmse_oos=0.15))
    assert (val, which) == (0.15, "out-of-sample")


def test_ashrae_falls_back_to_insample() -> None:
    val, which = verdicts.ashrae_cvrmse_used(make_model(cv_rmse=0.10, cv_rmse_oos=None))
    assert (val, which) == (0.10, "in-sample")


def test_ashrae_pass_boundary_inclusive() -> None:
    passed, _ = verdicts.ashrae_verdict(make_model(cv_rmse_oos=0.20))
    assert passed is True


def test_ashrae_fail_above_threshold() -> None:
    passed, text = verdicts.ashrae_verdict(make_model(cv_rmse_oos=0.21))
    assert passed is False
    assert "FAIL" in text


# --- F. Significance wording -------------------------------------------------
def test_significance_sentence_significant() -> None:
    sent = verdicts.significance_sentence(make_summary(is_significant=True))
    assert "significant" in sent.lower()
    assert "not statistically significant" not in sent.lower()


def test_significance_sentence_insignificant() -> None:
    sent = verdicts.significance_sentence(make_summary(is_significant=False, p_value=0.4))
    assert "not statistically significant" in sent.lower()
    assert "caution" in sent.lower()


def test_significance_sentence_nan() -> None:
    sent = verdicts.significance_sentence(make_summary(p_value=float("nan"), t_stat=float("nan")))
    assert "could not be tested" in sent.lower()


# --- G. Degenerate inputs ----------------------------------------------------
def test_insignificant_result_still_builds_both() -> None:
    s = make_summary(is_significant=False, p_value=0.4)
    assert build_mv_report_pdf(s, make_model(), make_avoided(), make_meta())[:5] == b"%PDF-"
    assert build_mv_report_xlsx(s, make_model(), make_avoided(), make_meta())[:4] == b"PK\x03\x04"


def test_nan_significance_fields_handled() -> None:
    s = make_summary(t_stat=float("nan"), p_value=float("nan"), lag1_autocorr=float("nan"),
                     n_effective=float("nan"))
    cells = _summary_cells(build_mv_report_xlsx(s, make_model(), make_avoided(), make_meta()), "Summary")
    assert cells["p-value"] == "n/a"
    assert build_mv_report_pdf(s, make_model(), make_avoided(), make_meta())[:5] == b"%PDF-"


def test_failing_model_injects_limitation_warning() -> None:
    text = _pdf_text(build_mv_report_pdf(make_summary(), make_model(cv_rmse_oos=0.35), make_avoided(), make_meta()))
    assert "ASHRAE" in text and ("not" in text.lower())


def test_empty_avoided_df_raises() -> None:
    empty = make_avoided(0)
    with pytest.raises(ValueError, match="empty"):
        build_mv_report_pdf(make_summary(), make_model(), empty, make_meta())
    with pytest.raises(ValueError, match="empty"):
        build_mv_report_xlsx(make_summary(), make_model(), empty, make_meta())


# --- I. No-Streamlit guarantee (mirrors engine purity) -----------------------
def test_reporting_imports_no_streamlit() -> None:
    code = (
        "import sys; import reporting, reporting.pdf, reporting.xlsx, reporting.charts, "
        "reporting.verdicts, reporting.format, reporting.meta; "
        "assert 'streamlit' not in sys.modules, 'reporting must not import streamlit'"
    )
    result = subprocess.run([sys.executable, "-c", code], cwd=str(ROOT), capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_reporting_sources_have_no_streamlit_import() -> None:
    for py in (ROOT / "reporting").glob("*.py"):
        assert "import streamlit" not in py.read_text(encoding="utf-8")
