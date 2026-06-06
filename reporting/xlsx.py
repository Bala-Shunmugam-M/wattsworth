"""Excel (.xlsx) M&V report — the machine-checkable, round-trippable artifact."""
from __future__ import annotations

import math
from io import BytesIO
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl import Workbook

from reporting import verdicts
from reporting.meta import ReportMeta

if TYPE_CHECKING:
    from engine.baseline import BaselineModel
    from engine.mv import SavingsSummary


def _na(value: object) -> object:
    """Render None/NaN as 'n/a' for cells; pass numbers through unchanged."""
    if value is None:
        return "n/a"
    if isinstance(value, float) and math.isnan(value):
        return "n/a"
    return value


def build_mv_report_xlsx(
    summary: "SavingsSummary",
    model: "BaselineModel",
    avoided_df: pd.DataFrame,
    meta: ReportMeta,
) -> bytes:
    """Build the .xlsx M&V report and return the file bytes.

    Raises:
        ValueError: If ``avoided_df`` is empty.
    """
    if avoided_df.empty:
        raise ValueError("Cannot build report: avoided_df is empty.")

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    for row in [
        ("Facility", meta.facility_name),
        ("Report ID", meta.report_id),
        ("Avoided energy (kWh)", summary.avoided_kwh),
        ("Avoided energy (MWh)", summary.avoided_kwh / 1000.0),
        ("Saving (%)", summary.pct_saving * 100.0),
        ("Cost saved (INR)", summary.inr_saved),
        ("CO2 avoided (t)", summary.tonnes_co2_avoided),
        ("Reporting days", summary.days),
        ("CI confidence", summary.ci_confidence),
        ("Avoided kWh CI low", _na(summary.avoided_kwh_ci_low)),
        ("Avoided kWh CI high", _na(summary.avoided_kwh_ci_high)),
        ("Annual kWh", summary.annual_kwh),
        ("Annual INR", summary.annual_inr),
        ("Annual tCO2", summary.annual_tco2),
        ("p-value", _na(summary.p_value)),
        ("Significant", bool(summary.is_significant)),
    ]:
        ws.append(row)

    wsb = wb.create_sheet("Baseline")
    _passed, verdict_text = verdicts.ashrae_verdict(model)
    for row in [
        ("R-squared", model.r_squared),
        ("CV(RMSE) in-sample", model.cv_rmse),
        ("CV(RMSE) out-of-sample", _na(model.cv_rmse_oos)),
        ("Durbin-Watson", model.durbin_watson),
        ("Max VIF", _na(model.max_vif)),
        ("Breusch-Pagan p", _na(model.bp_pvalue)),
        ("Observations", model.n_obs),
        ("Baseline period", str(model.baseline_period)),
        ("ASHRAE verdict", verdict_text),
    ]:
        wsb.append(row)
    wsb.append(("", ""))
    wsb.append(("Coefficient", "Value"))
    for term, coef in model.coefficients.items():
        wsb.append((term, float(coef)))

    wsa = wb.create_sheet("Avoided")
    wsa.append([str(c) for c in avoided_df.columns])
    for _, r in avoided_df.iterrows():
        wsa.append([(x.isoformat() if hasattr(x, "isoformat") else x) for x in r.tolist()])

    wsm = wb.create_sheet("Methodology")
    for row in [
        ("Tariff (INR/kWh)", meta.tariff_inr_per_kwh),
        ("Emission factor (kgCO2/kWh)", meta.emission_factor),
        ("IPMVP option", meta.ipmvp_option),
        ("Baseline method", "OLS + backward elimination of insignificant drivers"),
        ("Significance method", "autocorrelation-corrected one-sided t-test (effective N)"),
        ("Software", meta.software_version),
        ("Generated at", meta.generated_at),
    ]:
        wsm.append(row)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
